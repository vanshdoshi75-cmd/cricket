"""
excel_export.py
Builds a formatted .xlsx export of bookings using openpyxl.
"""

import os
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXPORT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")


def _date_range_for_filter(filter_name, start=None, end=None):
    today = date.today()
    if filter_name == "today":
        return today.isoformat(), today.isoformat()
    if filter_name == "week":
        start_d = today - timedelta(days=today.weekday())
        return start_d.isoformat(), today.isoformat()
    if filter_name == "month":
        start_d = today.replace(day=1)
        return start_d.isoformat(), today.isoformat()
    if filter_name == "custom" and start and end:
        return start, end
    return None, None  # all bookings


def export_bookings_to_excel(conn, filter_name="all", start=None, end=None):
    os.makedirs(EXPORT_DIR, exist_ok=True)

    query = "SELECT * FROM bookings"
    params = []
    date_start, date_end = _date_range_for_filter(filter_name, start, end)
    if date_start and date_end:
        query += " WHERE booking_date BETWEEN ? AND ?"
        params = [date_start, date_end]
    query += " ORDER BY booking_date DESC, start_time ASC"

    rows = conn.execute(query, params).fetchall()
    turf = conn.execute("SELECT name FROM turf LIMIT 1").fetchone()
    turf_name = turf["name"] if turf else "Cricket Turf"

    wb = Workbook()
    ws = wb.active
    ws.title = "Bookings"

    headers = [
        "Booking ID", "Customer Name", "Mobile", "Email", "Turf Name",
        "Booking Date", "Day", "Start Time", "End Time", "Number of Hours",
        "Price Per Hour", "Total Amount", "Booking Status", "Booking Created Date",
        "Payment Method", "Amount Paid", "Remaining Amount", "Payment Status",
        "Transaction ID", "Paid At",
    ]
    header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    money_columns = {12, 16, 17}  # Total Amount, Amount Paid, Remaining Amount

    for r_idx, b in enumerate(rows, start=2):
        row_keys = b.keys()
        payment_method = (b["payment_method"] if "payment_method" in row_keys else None) or "pay_at_turf"
        amount_paid = (b["amount_paid"] if "amount_paid" in row_keys else None) or 0
        remaining_amount = (b["remaining_amount"] if "remaining_amount" in row_keys else None)
        if remaining_amount is None:
            remaining_amount = b["total_amount"] - amount_paid
        payment_status = (b["payment_status"] if "payment_status" in row_keys else None) or "pay_at_turf"
        transaction_id = (b["transaction_id"] if "transaction_id" in row_keys else None) or "-"
        paid_at = (b["paid_at"] if "paid_at" in row_keys else None) or "-"

        values = [
            b["booking_code"], b["customer_name"], b["mobile"], b["email"] or "-",
            turf_name, b["booking_date"], b["day_name"], b["start_time"], b["end_time"],
            b["hours"], b["price_per_hour"], b["total_amount"], b["status"].title(),
            b["created_at"],
            payment_method.replace("_", " ").title(), amount_paid, remaining_amount,
            payment_status.replace("_", " ").title(), transaction_id, paid_at,
        ]
        for c_idx, v in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.border = border
            if c_idx in money_columns:
                cell.number_format = '"₹"#,##0'

    widths = [16, 20, 14, 24, 22, 14, 12, 11, 11, 12, 13, 13, 14, 20, 14, 13, 16, 20, 18, 18]
    for i, w in enumerate(widths, start=1):
        col_letter = chr(64 + i) if i <= 26 else "A" + chr(64 + (i - 26))
        ws.column_dimensions[col_letter].width = w

    ws.freeze_panes = "A2"

    filename = f"bookings_export_{date.today().isoformat()}_{filter_name}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)
    return filepath, filename, len(rows)
